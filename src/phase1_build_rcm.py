"""Phase 1 - build the Risk and Control Matrix workbook (outputs/rcm/SOX_ITGC_RCM.xlsx)."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

import common as C
from rcm_definitions import CONTROLS, RISKS, risk_by_id, severity

NAVY = "1F3864"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(bold=True, color="FFFFFF")
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
SEV_FILL = {"Critical": "F4B6B6", "High": "FAD7B5", "Medium": "FFF2B3", "Low": "D9EAD3"}


def write_table(ws, headers, rows, widths, name):
    ws.append(headers)
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment, cell.border = WRAP, THIN
    ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(t)
    ws.freeze_panes = "B2"


def build():
    wb = Workbook()

    # --- Cover ---
    ws = wb.active
    ws.title = "Cover"
    lines = [
        ("SOX ITGC Risk and Control Matrix", Font(bold=True, size=16, color=NAVY)),
        (f"Entity: {C.COMPANY}", None),
        (f"Period: {C.FISCAL_YEAR} ({C.PERIOD_START:%d %b %Y} - {C.PERIOD_END:%d %b %Y})", None),
        ("Framework: COSO 2013 Internal Control - Integrated Framework (Principle 11: general controls over technology)", None),
        ("", None),
        ("SYNTHETIC DATA NOTICE: This matrix, the entity, systems, people and all data are fictional and were "
         "created for this simulation. It does not reproduce any real company's RCM.", Font(bold=True, color="C00000")),
        ("", None),
        ("Scope", Font(bold=True, size=12)),
        ("In-scope applications: " + "; ".join(f"{k} - {v}" for k, v in C.SYSTEMS.items()), None),
        ("ITGC domains: User Access (provisioning / deprovisioning / review), Privileged Access, Change Management", None),
        ("Out of scope for this simulation: IT operations (job scheduling, backup), network/infrastructure, SOC 1 reliance", None),
        ("", None),
        ("How to read this workbook", Font(bold=True, size=12)),
        ("Risks - financial reporting risks with inherent likelihood x impact severity score (1-25)", None),
        ("RCM - one row per control: owner, frequency, type, nature, key-control flag, population and test attribute", None),
        ("Policy Parameters - thresholds each control is tested against in Phase 3", None),
        ("Severity scale: Critical >= 15, High 10-14, Medium 5-9, Low < 5", None),
    ]
    for text, font in lines:
        ws.append([text])
        if font:
            ws.cell(ws.max_row, 1).font = font
    ws.column_dimensions["A"].width = 130
    for row in ws.iter_rows():
        row[0].alignment = Alignment(wrap_text=True, vertical="top")

    # --- Risks ---
    ws = wb.create_sheet("Risks")
    rows = []
    for r in RISKS:
        score, rating = severity(r)
        n_controls = sum(c["risk_id"] == r["risk_id"] for c in CONTROLS)
        rows.append([r["risk_id"], r["title"], r["statement"], r["assertions"], r["fs_impact"],
                     r["likelihood"], r["impact"], score, rating, n_controls])
    write_table(ws, ["Risk ID", "Risk Title", "Risk Statement", "Relevant Assertions",
                     "Financial Statement Impact", "Likelihood (1-5)", "Impact (1-5)",
                     "Severity Score", "Inherent Rating", "# Mitigating Controls"],
                rows, [9, 30, 70, 24, 40, 11, 11, 10, 12, 12], "Risks")
    for row in ws.iter_rows(min_row=2):
        row[8].fill = PatternFill("solid", fgColor=SEV_FILL[row[8].value])

    # --- RCM ---
    ws = wb.create_sheet("RCM")
    rows = []
    for c in CONTROLS:
        r = risk_by_id(c["risk_id"])
        rows.append([c["control_id"], c["domain"], c["risk_id"], r["title"], c["title"],
                     c["description"], c["owner"], c["frequency"], c["type"], c["nature"],
                     c["key"], "; ".join(C.SYSTEMS), r["assertions"], "COSO P11 / P12",
                     c["population"], c["attribute"], c["procedure"]])
    write_table(ws, ["Control ID", "ITGC Domain", "Risk ID", "Risk Addressed", "Control Title",
                     "Control Description", "Control Owner", "Frequency", "Control Type",
                     "Control Nature", "Key Control", "Systems", "Assertions", "COSO Principle",
                     "Test Population (source)", "Test Attribute / Exception Definition",
                     "Test of Operating Effectiveness"],
                rows, [13, 17, 8, 28, 28, 60, 22, 18, 12, 18, 8, 20, 22, 12, 40, 40, 60], "RCM")
    ws.row_dimensions[1].height = 32

    # --- Policy parameters ---
    ws = wb.create_sheet("Policy Parameters")
    params = [
        ["Termination revocation SLA", f"{C.TERMINATION_REVOCATION_SLA_BD} business days", "ITGC-UA-02"],
        ["HR separation notification SLA", f"{C.HR_NOTIFICATION_SLA_BD} business day", "ITGC-UA-03"],
        ["Quarterly access review due", f"{C.ACCESS_REVIEW_DUE_DAYS} calendar days after quarter end", "ITGC-UA-04"],
        ["Privileged access levels", ", ".join(C.PRIVILEGED_ACCESS_LEVELS), "ITGC-PA-01 / PA-02"],
        ["Emergency change retrospective approval", f"{C.EMERGENCY_RETRO_APPROVAL_BD} business days after implementation", "ITGC-CM-01"],
        ["Business-day calendar", "Mon-Fri excluding NYSE holidays", "All SLA tests"],
        ["Exception aging as-of date", f"{C.TEST_AS_OF:%Y-%m-%d}", "Remediation prioritization"],
    ]
    write_table(ws, ["Parameter", "Value", "Used By"], params, [40, 55, 22], "Params")

    C.RCM_DIR.mkdir(parents=True, exist_ok=True)
    out = C.RCM_DIR / "SOX_ITGC_RCM.xlsx"
    wb.save(out)
    print(f"[Phase 1] RCM written: {out.relative_to(C.ROOT)} ({len(RISKS)} risks, {len(CONTROLS)} controls)")


if __name__ == "__main__":
    build()
